"""Build the dashboard dataset for Tatva Global School (TGS).

Inputs (pipeline/raw):
  meta_posts.json   per-post Meta Business Suite export (IG + FB, 90 days), with *_ads / *_org splits
  meta_items.json   the same posts with Meta content ids (for linking images)
  meta_img/         full post images grabbed from Meta post-insights pages ({content_id}.jpg)
  calendar.xlsx     the content calendar sheet (<Month> details + <Month> Snapshot tabs)
  oct_items.json    October plan, already cleaned
  canva/            creatives exported from Canva / the plan deck
  ../reconcile.json manual decisions (plan <-> post matches, markers, titles)

Output:
  app/src/data/tgs.json      one JSON document the app imports
  app/public/creatives/...   images referenced by that JSON
"""
import json, re, os, shutil, statistics as st, unicodedata, datetime as dt
from difflib import SequenceMatcher
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, 'raw')
APP = os.path.join(HERE, '..', 'app')
OUT_JSON = os.path.join(APP, 'src', 'data', 'tgs.json')
OUT_IMG = os.path.join(APP, 'public', 'creatives')
YEAR = 2026
REPORT_MONTHS = [7, 8, 9]
PLAN_MONTHS = [10]
MONTH_NAMES = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
               'September', 'October', 'November', 'December']

REC = json.load(open(os.path.join(HERE, 'reconcile.json')))
CUTOFF = dt.date.fromisoformat(REC['dataCutoff'])

METRICS = ['Views', 'Reach', 'Interactions', 'Likes and reactions', 'Comments', 'Shares', 'Saves',
           'Link clicks', 'Follows']
MKEY = {'Views': 'views', 'Reach': 'reach', 'Interactions': 'inter', 'Likes and reactions': 'likes',
        'Comments': 'comments', 'Shares': 'shares', 'Saves': 'saves', 'Link clicks': 'clicks',
        'Follows': 'follows'}


def nfkc(s):
    return unicodedata.normalize('NFKC', s or '')


def norm(s):
    s = nfkc(s).lower()
    s = re.sub(r'[^a-z0-9 ]+', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def num(v):
    if v is None:
        return 0
    if isinstance(v, str):  # raw cells like '16565 14734 from ads' -> total is the first number
        m = re.match(r'\s*([\d,.]+)', v)
        return float(m.group(1).replace(',', '')) if m else 0
    return v


def ads_of(v):
    if isinstance(v, str):
        m = re.search(r'([\d,.]+)\s+from ads', v)
        return float(m.group(1).replace(',', '')) if m else 0
    return 0


def secs(v):
    if not isinstance(v, str):
        return v
    t = 0
    for n, u in re.findall(r'(\d+)\s*([dhms])', v):
        t += int(n) * {'d': 86400, 'h': 3600, 'm': 60, 's': 1}[u]
    return t


def ts(d):
    return d.strftime('%Y-%m-%d %H:%M')


# --------------------------------------------------------------------------- Meta posts
posts = json.load(open(os.path.join(RAW, 'meta_posts.json')))
items = json.load(open(os.path.join(RAW, 'meta_items.json')))['ITEMS']
for p in posts:
    p['_d'] = dt.datetime.strptime(p['date'], '%Y-%m-%d %H:%M')
    p['_story'] = p['type'] == 'Story'


def item_dt(s):
    day, mon, hm = s.split(' ')
    h, m = hm.split(':')
    return dt.datetime(YEAR, MONTH_NAMES.index(mon) + 1, int(day), int(h), int(m))


for it in items:
    it['_d'] = item_dt(it['date'])
    it['_story'] = 'STORY' in it['type']
    it['_ig'] = it['type'].startswith('IG')
    it['_used'] = False

# join each post record to its Meta id
unjoined = []
for p in posts:
    want_ig = p['platform'] == 'IG'
    cands = [it for it in items if not it['_used'] and it['_d'] == p['_d'] and it['_story'] == p['_story']
             and it['_ig'] == want_ig]
    if not cands:
        cands = [it for it in items if not it['_used'] and it['_d'] == p['_d'] and it['_story'] == p['_story']]
    if not cands:
        unjoined.append(p)
        continue
    best = max(cands, key=lambda it: SequenceMatcher(None, norm(it['cap'])[:40], norm(p['caption'])[:40]).ratio())
    best['_used'] = True
    p['id'] = best['id']
assert not unjoined, f'{len(unjoined)} posts without Meta id'


def family(acc):
    a = (acc or '').lower()
    if 'jumpytalks' in a:
        return 'jumpytalks'
    if 'basketball' in a:
        return 'tatva_basketball_academy_'
    return 'tgs'


def fmt_of(t):
    return {'Reel': 'Reel', 'Photo': 'Static', 'Carousel': 'Carousel', 'Multi media': 'Carousel'}.get(t)


# --------------------------------------------------------------------------- group IG/FB copies into content
feed = sorted([p for p in posts if not p['_story']], key=lambda p: p['_d'])
groups = []
for p in feed:
    fam = family(p['account'])
    key = norm(p['caption'])[:60]
    home = None
    for g in groups:
        if g['family'] != fam:
            continue
        if abs((g['recs'][0]['_d'] - p['_d']).total_seconds()) > 36 * 3600:
            continue
        if any(r['platform'] == p['platform'] or r['platform'] == 'IG+FB' for r in g['recs']) or p['platform'] == 'IG+FB':
            continue
        gk = norm(g['recs'][0]['caption'])[:60]
        close = abs((g['recs'][0]['_d'] - p['_d']).total_seconds()) <= 300
        sim = SequenceMatcher(None, gk, key).ratio() if key else 0
        if key and (gk == key or sim > 0.9 or (close and sim > 0.6)):
            home = g
            break
    if home:
        home['recs'].append(p)
    else:
        groups.append({'family': fam, 'recs': [p]})

img_meta = json.load(open(os.path.join(RAW, 'meta_img', '_meta.json')))
os.makedirs(os.path.join(OUT_IMG, 'm'), exist_ok=True)
os.makedirs(os.path.join(OUT_IMG, 'p'), exist_ok=True)


def copy_meta_img(ids):
    for i in ids:
        src = os.path.join(RAW, 'meta_img', f'{i}.jpg')
        if os.path.exists(src) and i in img_meta:
            from PIL import Image, ImageStat
            if ImageStat.Stat(Image.open(src).convert('L').resize((64, 80))).stddev[0] < 15:
                return None  # blank frame (e.g. a reel's plain intro) is not a preview of the post
            shutil.copy(src, os.path.join(OUT_IMG, 'm', f'{i}.jpg'))
            m = img_meta[i]
            return {'src': f'creatives/m/{i}.jpg', 'w': m['w'], 'h': m['h'], 'from': 'meta'}
    return None


content = []
for g in groups:
    recs = sorted(g['recs'], key=lambda r: (0 if r['platform'] == 'IG' else 1, r['_d']))
    first = min(r['_d'] for r in recs)
    plats = set()
    for r in recs:
        plats |= set(r['platform'].split('+'))
    ig = [r for r in recs if r['platform'] == 'IG']
    cid = (ig[0] if ig else recs[0])['id']
    fmt = next((fmt_of(r['type']) for r in recs if fmt_of(r['type'])), None)
    m = {}
    for k in METRICS:
        tot = sum(num(r.get(k)) for r in recs)
        ads = sum(ads_of(r.get(k)) for r in recs)
        m[MKEY[k]], m[MKEY[k] + '_ads'], m[MKEY[k] + '_org'] = tot, ads, tot - ads
    plays = [secs(r.get('Video average play time')) for r in recs if r.get('Video average play time')]
    m['avgPlay'] = max(plays) if plays else None
    per = {}
    for r in recs:
        for pl in r['platform'].split('+'):
            per.setdefault(pl, {'views': 0, 'reach': 0, 'inter': 0})
        tgt = r['platform'] if r['platform'] in ('IG', 'FB') else 'IG+FB'
        per.setdefault(tgt, {'views': 0, 'reach': 0, 'inter': 0})
        per[tgt]['views'] += num(r.get('Views'))
        per[tgt]['reach'] += num(r.get('Reach'))
        per[tgt]['inter'] += num(r.get('Interactions'))
    per = {k: v for k, v in per.items() if v['views'] or v['reach']}
    caption = nfkc(max((r['caption'] for r in recs), key=len))
    img = copy_meta_img([cid] + [r['id'] for r in recs])
    content.append({
        'id': cid, 'ids': [r['id'] for r in recs], 'family': g['family'],
        'account': recs[0]['account'], 'date': ts(first), 'month': f'{first.year}-{first.month:02d}',
        'platforms': 'IG+FB' if plats >= {'IG', 'FB'} else ''.join(sorted(plats)),
        'crossposted': any(r['platform'] == 'IG+FB' for r in recs),
        'format': fmt, 'caption': caption, 'metrics': m, 'perPlatform': per, 'image': img,
        '_stamps': [ts(r['_d']) for r in recs],
    })
for c in content:
    c['metrics'] = {k: (int(v) if isinstance(v, float) and v.is_integer() else v) for k, v in c['metrics'].items()}
    for pv in c['perPlatform'].values():
        for k in pv:
            pv[k] = int(pv[k])
content.sort(key=lambda c: c['date'])
by_first = {c['date']: c for c in content}
by_any_ts = {}
for c in content:
    for stamp in c.pop('_stamps'):
        by_any_ts.setdefault(stamp, c)


def find_content(stamp):
    c = by_first.get(stamp) or by_any_ts.get(stamp)
    if not c:
        raise SystemExit(f'reconcile.json points at {stamp} but no post has that timestamp')
    return c


# --------------------------------------------------------------------------- stories
stories = {}
for p in posts:
    if p['_story']:
        k = f"{p['_d'].year}-{p['_d'].month:02d}"
        s = stories.setdefault(k, {'count': 0, 'views': 0})
        s['count'] += 1
        s['views'] += int(num(p.get('Views')))

# --------------------------------------------------------------------------- calendar
wb = openpyxl.load_workbook(os.path.join(RAW, 'calendar.xlsx'), data_only=True)
DETAIL_TAB = {7: 'July Details', 8: 'August details', 9: 'September details', 10: 'October details'}
SNAP_TAB = {7: 'July Snapshot', 8: 'August Snapshot', 9: 'September Snapshot', 10: 'October Snapshot'}


def snapshot(month):
    """date -> {platforms, label, boost, video} from the calendar grid (date row, platform row, label row)."""
    ws = wb[SNAP_TAB[month]]
    out = {}
    for r in range(1, min(ws.max_row, 60) + 1):
        for c in range(1, 8):
            v = ws.cell(r, c).value
            if isinstance(v, dt.datetime) and v.month == month:
                plat = ws.cell(r + 1, c).value
                lab = ws.cell(r + 2, c).value
                lab = str(lab or '')
                out[v.date()] = {
                    'platforms': 'LinkedIn + Meta' if plat and 'Ln' in str(plat) else ('Meta' if plat else None),
                    'label': lab.replace('💵', '').replace('📹', '').replace('\n', ' ').strip(' /'),
                    'boost': '💵' in lab, 'video': '📹' in lab}
    return out


def details(month):
    ws = wb[DETAIL_TAB[month]]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hdr = [(str(c).replace('\n', ' ').strip().lower() if c else '') for c in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(hdr):
                if h.startswith(n):
                    return i
        return None
    ix = dict(designer=col('designer'), status=col('status'), req=col('requirement'), date=col('date'),
              bucket=col('bucket'), fmt=col('format'), concept=col('post concept'), copy=col('creative copy'),
              caption=col('caption'), visual=col('visual ideas'))
    out, seen = [], {}
    for r in rows[1:]:
        g = lambda k: (r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None)
        d = g('date')
        if isinstance(d, str):  # e.g. 'Jul 16 (Thu)\n Wk 3 — Q2 Boost'
            mm = re.match(r'\s*([A-Za-z]{3})[a-z]*\.?\s+(\d{1,2})', d)
            if mm and mm.group(1).title() in [n[:3] for n in MONTH_NAMES]:
                d = dt.datetime(YEAR, [n[:3] for n in MONTH_NAMES].index(mm.group(1).title()) + 1, int(mm.group(2)))
        if not isinstance(d, dt.datetime):
            continue
        n = seen.get(d.date(), 0)
        seen[d.date()] = n + 1
        clean = lambda v: re.sub(r'\s+\n', '\n', str(v).strip()) if v not in (None, '') else ''
        out.append(dict(key=f'{d.date()}#{n}', date=d.date(), designer=clean(g('designer')),
                        sheetStatus=clean(g('status')), req=clean(g('req')).lower(),
                        bucket=re.sub(r'\s+', ' ', clean(g('bucket'))), fmt=clean(g('fmt')),
                        concept=clean(g('concept')), copy=clean(g('copy')), caption=clean(g('caption')),
                        visual=clean(g('visual'))))
    return out


def plan_format(f, bucket):
    f = (f or '').lower()
    if 'reel' in f:
        return 'Reel'
    if 'carousel' in f:
        return 'Carousel'
    if 'static' in f:
        return 'Static'
    return 'Reel' if 'highlight' in (bucket or '').lower() else 'Static'


def owner_of(req, fmt):
    if req.startswith('zam'):
        return 'Zamstars', False
    if req.startswith('tatva'):
        return 'Tatva', False
    return ('Tatva' if fmt == 'Reel' else 'Zamstars'), True


def dependency(owner, fmt, copy=''):
    if 'shared by tatva' in (copy or '').lower() or (fmt == 'Reel' and owner == 'Tatva'):
        return 'Footage from Tatva'
    if fmt == 'Reel':
        return None
    if owner == 'Tatva':
        return 'Photo from Tatva'
    return None


os.makedirs(os.path.join(OUT_IMG, 'p'), exist_ok=True)


def plan_creative(fname):
    if not fname:
        return None
    src = os.path.join(RAW, 'canva', fname)
    if not os.path.exists(src):
        return None
    shutil.copy(src, os.path.join(OUT_IMG, 'p', fname))
    from PIL import Image
    w, h = Image.open(src).size
    return {'src': f'creatives/p/{fname}', 'w': w, 'h': h, 'from': 'canva'}


plan = []
assigned = set()
for month in REPORT_MONTHS:
    snap = snapshot(month)
    for row in details(month):
        rec = REC['plan'].get(row['key'], {})
        fmt = plan_format(row['fmt'], row['bucket'])
        owner, inferred = owner_of(row['req'], fmt)
        sn = snap.get(row['date'], {})
        match = None
        if 'live' in rec:
            if rec['live']:
                match = find_content(rec['live'])
        else:  # automatic: caption similarity within 4 days
            best = None
            for c in content:
                if c['family'] != 'tgs' or c['id'] in assigned:
                    continue
                d = dt.datetime.strptime(c['date'], '%Y-%m-%d %H:%M').date()
                dd = abs((d - row['date']).days)
                if dd > 4:
                    continue
                a, b = norm(row['caption'])[:180], norm(c['caption'])[:180]
                sim = SequenceMatcher(None, a, b).ratio() if a and b else 0
                score = sim + (0.25 if dd == 0 else 0.1 if dd <= 1 else 0)
                if best is None or score > best[0]:
                    best = (score, sim, dd, c)
            if best and (best[1] > 0.45 or (best[2] == 0 and best[1] > 0.2)):
                match = best[3]
            else:
                raise SystemExit(f"No automatic match for {row['key']} ({row['bucket']}); add it to reconcile.json")
        if match:
            assert match['id'] not in assigned, f"{match['id']} matched twice"
            assigned.add(match['id'])
        live_d = dt.datetime.strptime(match['date'], '%Y-%m-%d %H:%M').date() if match else None
        if match:
            lag = (live_d - row['date']).days
            status = 'On date' if lag == 0 else ('Late' if lag > 0 else 'Early')
        elif rec.get('replacedBy'):
            lag, status = None, 'Replaced'
        elif row['date'] > CUTOFF:
            lag, status = None, 'Upcoming'
        else:
            lag, status = None, 'Not posted'
        plan.append({
            'id': row['key'], 'month': f"{YEAR}-{month:02d}", 'date': str(row['date']),
            'title': rec.get('title') or row['bucket'], 'bucket': row['bucket'], 'format': fmt,
            'platforms': sn.get('platforms') or 'Meta', 'owner': owner, 'ownerInferred': inferred,
            'dependency': dependency(owner, fmt, row['copy']), 'designer': row['designer'] or None,
            'sheetStatus': row['sheetStatus'] or None, 'boostPlanned': bool(sn.get('boost')),
            'concept': row['concept'] or None, 'copy': row['copy'] or None, 'caption': row['caption'] or None,
            'status': status, 'lag': lag, 'contentId': match['id'] if match else None,
            'replacedBy': find_content(rec['replacedBy'])['id'] if rec.get('replacedBy') else None,
            'confidence': rec.get('confidence', 'confirmed' if match else None),
            'note': rec.get('note'), 'creative': plan_creative(rec.get('creative')),
            'needsFootage': False, 'footageDue': None, 'dueText': None, 'need': None,
        })

# October: already-cleaned plan items + snapshot boost flags
oct_snap = snapshot(10)
oct_rows = {str(r['date']): r for r in details(10)}
seen_undated = 0
for o in json.load(open(os.path.join(RAW, 'oct_items.json'))):
    d = o.get('date')
    sn = oct_snap.get(dt.date.fromisoformat(d), {}) if d else {}
    key = f'{d}#0' if d else f'undated#{seen_undated}'
    if not d:
        seen_undated += 1
    fmt = o['format']
    owner = o['source']
    img = o.get('image')
    plan.append({
        'id': key, 'month': '2026-10', 'date': d, 'title': o['title'], 'bucket': o['bucket'], 'format': fmt,
        'platforms': o.get('platforms') or 'Meta', 'owner': owner, 'ownerInferred': False,
        'dependency': o.get('sourceLabel') if o.get('sourceLabel') != 'Made by Zamstars' else None,
        'designer': o.get('designer') or None, 'sheetStatus': o.get('sheetStatus') or None,
        'boostPlanned': bool(sn.get('boost')), 'concept': o.get('concept'), 'copy': o.get('copy'),
        'caption': o.get('caption'), 'visual': o.get('visual'), 'state': o.get('state'),
        'status': 'Planned', 'lag': None, 'contentId': None, 'replacedBy': None, 'confidence': None,
        'note': o.get('note') or None,
        'creative': plan_creative(os.path.basename(img)) if img else None,
        'needsFootage': bool(o.get('needsFootage')), 'footageDue': o.get('footageDue'),
        'dueText': o.get('dueText'), 'need': o.get('need') if o.get('needsFootage') else None,
    })

# --------------------------------------------------------------------------- markers, titles, baselines
plan_by_content = {p['contentId']: p for p in plan if p['contentId']}
replaced_by = {p['replacedBy']: p for p in plan if p['replacedBy']}


def short_title(caption):
    s = nfkc(caption).strip()
    s = re.split(r'(?<=[.!?])\s', s)[0]
    s = s.split('\n')[0]
    return (s[:58].rsplit(' ', 1)[0] + '…') if len(s) > 60 else s


for c in content:
    p = plan_by_content.get(c['id'])
    un = REC['unplanned'].get(c['date'], {})
    if p:
        c['marker'] = 'planned'
        c['planId'] = p['id']
        c['title'] = p['title']
        c['owner'] = p['owner']
        if not c['format']:
            c['format'] = p['format']
    elif c['family'] != 'tgs':
        c['marker'] = 'collab'
        c['planId'] = None
        c['partner'] = REC['collabTitles'].get(c['family'], c['family'])
        c['title'] = un.get('title') or short_title(c['caption'])
        c['owner'] = None
    else:
        c['marker'] = un.get('marker', 'unconfirmed')
        c['planId'] = None
        c['title'] = un.get('title') or short_title(c['caption'])
        c['owner'] = 'Zamstars' if c['marker'] == 'lastminute' else None
    c['format'] = c['format'] or 'Static'
    c['evidence'] = un.get('evidence')
    c['note'] = un.get('note')
    c['replacesPlanId'] = replaced_by[c['id']]['id'] if c['id'] in replaced_by else None
    m = c['metrics']
    c['boosted'] = (m['views_ads'] > 0) or (m['reach_ads'] > 0)
    c['boostPlanned'] = bool(p and p['boostPlanned'])
    if not c['image'] and p and p.get('creative'):
        c['image'] = p['creative']

in_range = [c for c in content if int(c['month'][5:]) in REPORT_MONTHS and c['month'].startswith(str(YEAR))]
own = [c for c in in_range if c['family'] == 'tgs']
baselines = {}
for f in ('Reel', 'Static', 'Carousel'):
    vals = [c['metrics']['views_org'] for c in own if c['format'] == f]
    if vals:
        baselines[f] = {'n': len(vals), 'median': st.median(vals)}
for c in content:
    b = baselines.get(c['format'])
    c['index'] = round(c['metrics']['views_org'] / b['median'], 2) if b and c['family'] == 'tgs' else None

content = [c for c in content if c in in_range]  # June 30 falls outside the reporting window

# --------------------------------------------------------------------------- client view: plain-language numbers + verdict
FMT_WORD = {'Static': 'static post', 'Carousel': 'carousel', 'Reel': 'reel'}


def plain(c):
    m = c['metrics']
    return {'reached': m['reach_org'], 'reactions': m['likes_org'] + m['comments_org'],
            'passedOn': m['shares_org'] + m['saves_org'], 'follows': m['follows']}


own_rng = [c for c in content if c['family'] == 'tgs']
typical = {}
for f in ('Static', 'Carousel', 'Reel'):
    xs = [plain(c) for c in own_rng if c['format'] == f and not c['boosted']]
    typical[f] = {k: st.median([x[k] for x in xs]) for k in ('reached', 'reactions', 'passedOn', 'follows')}
    typical[f]['n'] = len(xs)


def ratio(a, b):
    return a / b if b else None


def fmt_k(x):
    return f"{x/1000:.1f}K".replace('.0K', 'K') if x >= 10000 else f"{int(round(x)):,}"


def times(k):
    return f"{k:.1f}×" if k < 10 else f"{k:.0f}×"


for c in content:
    p = plain(c)
    c['plain'] = p
    if c['family'] != 'tgs':
        c['verdict'] = None
        c['why'] = f"Posted by {c['partner']} with Tatva as collaborator, so it reached their audience."
        continue
    t = typical[c['format']]
    word = FMT_WORD[c['format']]
    m = c['metrics']
    if c['boosted']:
        c['verdict'] = {'key': 'boosted', 'label': 'Boosted', 'score': None}
        tot = m['reach']
        c['why'] = (f"Paid boost: {fmt_k(m['reach_ads'])} of the {fmt_k(tot)} people reached came through the ad, "
                    f"so it's judged on what the boost added rather than against unpaid posts.") if m['reach_ads'] > m['reach_org'] else \
                   (f"Boosted: {fmt_k(m['reach_ads'])} extra people reached through the ad, on top of {fmt_k(m['reach_org'])} organically.")
        continue
    ri = ratio(p['reached'], t['reached']) or 0
    eng = p['reactions'] + p['passedOn']
    ei = ratio(eng, t['reactions'] + t['passedOn']) or 0
    pi = ratio(p['passedOn'], t['passedOn'])
    score = (max(ri, 0.01) * max(ei, 0.01)) ** 0.5
    if score >= 1.5:
        v = {'key': 'standout', 'label': 'Standout'}
    elif score >= 0.75:
        v = {'key': 'steady', 'label': 'Steady'}
    else:
        v = {'key': 'quiet', 'label': 'Quiet'}
    v['score'] = round(score, 2)
    c['verdict'] = v
    if v['key'] == 'standout':
        drivers = [(ri, f"reached {times(ri)} as many people as a typical {word}", f"reached {times(ri)} the usual audience"),
                   (ei, f"got {times(ei)} the engagement of a typical {word}", f"got {times(ei)} the usual engagement")]
        if pi and p['passedOn'] >= 10:
            drivers.append((pi, f"was shared or saved {p['passedOn']:,} times, {times(pi)} a typical {word}",
                            f"was shared or saved {p['passedOn']:,} times ({times(pi)} usual)"))
        drivers.sort(key=lambda d: -d[0])
        why = f"It {drivers[0][1]}"
        if len(drivers) > 1 and drivers[1][0] >= 1.5:
            why += f", and {drivers[1][2]}"
        why += '.'
    elif v['key'] == 'steady':
        if ri < 0.85 and ei > 1.15:
            why = f"Reached fewer people than a typical {word} ({fmt_k(p['reached'])} vs {fmt_k(t['reached'])}), but those who saw it engaged more than usual."
        elif ri > 1.15 and ei < 0.85:
            why = f"Reached more people than a typical {word} ({fmt_k(p['reached'])} vs {fmt_k(t['reached'])}), but fewer of them reacted."
        elif ri >= 1.15:
            why = f"Reached more people than a typical {word} ({fmt_k(p['reached'])} vs {fmt_k(t['reached'])}); engagement was about usual."
        elif ri <= 0.85:
            why = f"Reached a few fewer people than a typical {word} ({fmt_k(p['reached'])} vs {fmt_k(t['reached'])}); engagement was about usual."
        else:
            why = f"In line with a typical {word}: {fmt_k(p['reached'])} people reached (usually {fmt_k(t['reached'])})."
    else:
        why = f"Reached {round(ri*100)}% of a typical {word}'s audience ({fmt_k(p['reached'])} vs {fmt_k(t['reached'])})."
    if p['follows'] >= 10:
        why += f" Brought in {p['follows']} new followers."
    c['why'] = why

# --------------------------------------------------------------------------- month summaries + insights


def fmt_n(x):
    x = round(x)
    return f'{x/1000:.1f}K'.replace('.0K', 'K') if x >= 10000 else f'{x:,}'


def month_summary(mk):
    pl = [p for p in plan if p['month'] == mk]
    cs = [c for c in content if c['month'] == mk]
    ow = [c for c in cs if c['family'] == 'tgs']
    col = [c for c in cs if c['family'] != 'tgs']
    status = {}
    for p in pl:
        status[p['status']] = status.get(p['status'], 0) + 1
    due = [p for p in pl if p['status'] != 'Upcoming']
    delivered = [p for p in due if p['contentId']]
    by_owner = {}
    for o in ('Zamstars', 'Tatva'):
        d = [p for p in due if p['owner'] == o]
        by_owner[o] = {'due': len(d), 'delivered': sum(1 for p in d if p['contentId'])}
    S = lambda arr, k: sum(c['metrics'][k] for c in arr)
    markers = {}
    for c in ow:
        markers[c['marker']] = markers.get(c['marker'], 0) + 1
    fmts = {}
    for c in ow:
        f = fmts.setdefault(c['format'], {'count': 0, 'views_org': 0, 'median': 0})
        f['count'] += 1
        f['views_org'] += c['metrics']['views_org']
    for f in fmts:
        fmts[f]['median'] = st.median([c['metrics']['views_org'] for c in ow if c['format'] == f])
    partners = {}
    for c in col:
        pz = partners.setdefault(c['partner'], {'count': 0, 'views': 0, 'inter': 0})
        pz['count'] += 1
        pz['views'] += c['metrics']['views']
        pz['inter'] += c['metrics']['inter']
    return {
        'planned': len(pl), 'due': len(due), 'delivered': len(delivered), 'status': status, 'byOwner': by_owner,
        'posts': len(ow), 'markers': markers, 'formats': fmts,
        'views_org': S(ow, 'views_org'), 'views_ads': S(ow, 'views_ads'), 'reach_org': S(ow, 'reach_org'),
        'inter_org': S(ow, 'inter_org'), 'shares': S(ow, 'shares'), 'saves': S(ow, 'saves'),
        'comments': S(ow, 'comments'), 'follows': S(ow, 'follows'), 'clicks': S(ow, 'clicks'),
        'boosted': sum(1 for c in ow if c['boosted']),
        'unplannedViewsShare': (S([c for c in ow if c['marker'] != 'planned'], 'views_org') / S(ow, 'views_org')) if ow else 0,
        'collab': {'posts': len(col), 'views': S(col, 'views'), 'partners': partners},
        'stories': stories.get(mk, {'count': 0, 'views': 0}),
        'plain': {k: sum(c['plain'][k] for c in ow) for k in ('reached', 'reactions', 'passedOn', 'follows')},
        'boostReach': S(ow, 'reach_ads'),
        'verdicts': {k: sum(1 for c in ow if c['verdict'] and c['verdict']['key'] == k) for k in ('standout', 'steady', 'quiet', 'boosted')},
    }


def insights(mk, summ, prev):
    out = []
    pl = [p for p in plan if p['month'] == mk]
    ow = [c for c in content if c['month'] == mk and c['family'] == 'tgs']
    # 1. delivery
    late = [p for p in pl if p['status'] == 'Late']
    miss = [p for p in pl if p['status'] == 'Not posted']
    tat_miss = [p for p in miss if p['dependency'] == 'Footage from Tatva']
    body = f"{summ['delivered']} of {summ['due']} planned posts went live"
    if late:
        body += f", {len(late)} of them late (by {min(p['lag'] for p in late)}–{max(p['lag'] for p in late)} days)" if len(late) > 1 else f", 1 late by {late[0]['lag']} days"
    body += '.'
    if miss:
        body += f" {len(miss)} didn't go out"
        if tat_miss:
            body += f"; {len(tat_miss)} of those {'was a reel' if len(tat_miss) == 1 else 'were reels'} waiting on footage from Tatva"
        body += '.'
    out.append({'kind': 'delivery', 'tone': 'neutral' if not miss else 'warn', 'title': 'Plan delivery', 'body': body,
                'refs': [p['id'] for p in miss]})
    # 2. strongest vs baseline
    ranked = sorted([c for c in ow if c['index'] is not None], key=lambda c: -c['index'])
    if ranked:
        t = ranked[0]
        out.append({'kind': 'top', 'tone': 'good', 'title': 'Strongest post against its format',
                    'body': f"“{t['title']}” got {fmt_n(t['metrics']['views_org'])} organic views, {t['index']:.1f}× the typical {t['format'].lower()} ({fmt_n(baselines[t['format']]['median'])}).",
                    'refs': [t['id']]})
    # 3. unplanned share
    unp = [c for c in ow if c['marker'] != 'planned']
    if unp and summ['views_org']:
        share = sum(c['metrics']['views_org'] for c in unp) / summ['views_org']
        lm = [c for c in unp if c['marker'] == 'lastminute']
        body = f"{len(unp)} post{'s' if len(unp) != 1 else ''} outside the calendar drew {share*100:.0f}% of organic views."
        if lm:
            body += f" {len(lm)} {'was a Zamstars last-minute post' if len(lm) == 1 else 'were Zamstars last-minute posts'}."
        out.append({'kind': 'unplanned', 'tone': 'neutral', 'title': 'Outside the plan', 'body': body,
                    'refs': [c['id'] for c in unp]})
    # 4. boosts
    bo = [c for c in ow if c['boosted']]
    planned_not = [p for p in pl if p['boostPlanned'] and p['contentId'] and
                   not next(c for c in content if c['id'] == p['contentId'])['boosted']]
    if bo or planned_not:
        parts = []
        for c in sorted(bo, key=lambda c: -max(c['metrics']['views_ads'], c['metrics']['reach_ads']))[:3]:
            m = c['metrics']
            if m['views_ads']:
                parts.append(f"“{c['title']}”: {fmt_n(m['views_ads'])} paid vs {fmt_n(m['views_org'])} organic views")
            else:
                parts.append(f"“{c['title']}”: {fmt_n(m['reach_ads'])} of {fmt_n(m['reach'])} reach from ads")
        body = (f"{len(bo)} boosted. " + '; '.join(parts) + (f'; +{len(bo) - 3} more' if len(bo) > 3 else '') + '.') if parts else 'Nothing was boosted.'
        if planned_not:
            body += f" {len(planned_not)} post{'s' if len(planned_not) > 1 else ''} marked 💵 in the calendar went live without a boost."
        out.append({'kind': 'boost', 'tone': 'neutral', 'title': 'Paid support', 'body': body,
                    'refs': [c['id'] for c in bo] + [p['id'] for p in planned_not]})
    # 5. month over month
    if prev:
        dv = summ['views_org'] - prev['views_org']
        df = summ['follows'] - prev['follows']
        out.append({'kind': 'trend', 'tone': 'good' if dv >= 0 else 'warn', 'title': 'Against last month',
                    'body': f"Organic views {'up' if dv >= 0 else 'down'} {abs(dv)/prev['views_org']*100:.0f}% ({fmt_n(prev['views_org'])} → {fmt_n(summ['views_org'])}). "
                            f"Follows from posts: {summ['follows']:,} vs {prev['follows']:,}.",
                    'refs': []})
    return out


months = []
prev = None
for mnum in REPORT_MONTHS:
    mk = f'{YEAR}-{mnum:02d}'
    summ = month_summary(mk)
    months.append({'id': mk, 'label': MONTH_NAMES[mnum - 1], 'year': YEAR, 'kind': 'report', 'summary': summ,
                   'insights': insights(mk, summ, prev)})
    prev = summ
for mnum in PLAN_MONTHS:
    mk = f'{YEAR}-{mnum:02d}'
    pl = [p for p in plan if p['month'] == mk]
    months.append({'id': mk, 'label': MONTH_NAMES[mnum - 1], 'year': YEAR, 'kind': 'plan', 'summary': {
        'planned': len(pl), 'dated': sum(1 for p in pl if p['date']),
        'designReady': sum(1 for p in pl if (p.get('state') or '').startswith('Design ready')),
        'needsFootage': sum(1 for p in pl if p['needsFootage']),
        'linkedin': sum(1 for p in pl if 'LinkedIn' in (p['platforms'] or '')),
        'boostPlanned': sum(1 for p in pl if p['boostPlanned']),
        'formats': {f: sum(1 for p in pl if p['format'] == f) for f in ('Static', 'Reel', 'Carousel')},
    }, 'insights': []})

data = {
    'generatedAt': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
    'dataCutoff': REC['dataCutoff'],
    'brand': {'id': 'tgs', 'name': 'Tatva Global School', 'short': 'TGS', 'ig': 'tatva_global_school',
              'fb': 'Tatva Global School'},
    'brands': [
        {'id': 'tgs', 'name': 'Tatva Global School', 'short': 'TGS', 'connected': True},
        {'id': 'tk', 'name': 'Tatva Kids', 'short': 'TK', 'connected': False,
         'reason': 'The Meta login used for this pilot has no access to the Tatva Kids page.'},
    ],
    'sources': [
        {'name': 'Meta Business Suite', 'detail': 'Per-post insights, last 90 days, IG + FB, organic/paid split'},
        {'name': 'Content calendar', 'detail': 'Google Sheet: <Month> details + Snapshot tabs'},
        {'name': 'Canva', 'detail': 'TGS monthly design files (creatives, last-minute evidence)'},
    ],
    'account': json.load(open(os.path.join(RAW, 'account.json'))),
    'baselines': baselines, 'typical': typical, 'months': months, 'plan': plan, 'content': content, 'stories': stories,
}
os.makedirs(os.path.dirname(OUT_JSON), exist_ok=True)
json.dump(data, open(OUT_JSON, 'w'), ensure_ascii=False, indent=1)

# --------------------------------------------------------------------------- console report
print(f"content {len(content)} (own {sum(1 for c in content if c['family']=='tgs')}, collab {sum(1 for c in content if c['family']!='tgs')}), plan {len(plan)}")
print('baselines', {k: (v['n'], v['median']) for k, v in baselines.items()})
for m in months:
    s = m['summary']
    if m['kind'] == 'report':
        print(m['label'], 'org', round(s['views_org']), 'ads', round(s['views_ads']), 'follows', s['follows'], 'inter', s['inter_org'],
              'posts', s['posts'], s['markers'], 'delivered', f"{s['delivered']}/{s['due']}", s['status'], 'stories', s['stories'])
missing = [c['id'] for c in content if not c['image']]
print('no image:', len(missing), missing)
